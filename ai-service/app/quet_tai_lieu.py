"""
quet_tai_lieu.py - Bước 1 và 2 của pipeline xử lý tài liệu.

  python quet_tai_lieu.py quet        # phân loại mọi file -> kiem_tra_file.csv
  python quet_tai_lieu.py chuyen-doc  # đổi .doc -> .docx bằng LibreOffice

Cài thư viện:  pip install pymupdf python-docx
Chuyển .doc cần LibreOffice: https://www.libreoffice.org/download/
Script KHÔNG sửa, KHÔNG xoá file gốc. Chưa được chạy thử trên dữ liệu thật của bạn.
"""
import argparse
import csv
import re
import shutil
import subprocess
import unicodedata
from pathlib import Path

NGUON_MAC_DINH = Path("ai-service/data/raw_documents")
RA_MAC_DINH = Path("ai-service/data/kiem_tra_file.csv")
DICH_DOC_MAC_DINH = Path("ai-service/data/doc_da_chuyen")

# Từ phổ biến trong văn bản hành chính tiếng Việt. Dùng để phát hiện lỗi font:
# font cũ (VNI/TCVN3) cho ra ký tự Latin-1 "hợp lệ" nhưng không ghép thành từ tiếng Việt.
TU_PHO_BIEN = ["của", "và", "điều", "sinh viên", "quy định", "học", "được",
               "trường", "các", "không", "theo", "trong", "đào tạo"]
NGUONG_TRANG_CO_CHU = 20      # trang có ít hơn số ký tự này coi là trang ảnh
NGUONG_TI_LE_HOP_LE = 0.85

COT = ["file", "dinh_dang", "kb", "so_trang", "trang_co_chu", "so_ky_tu",
       "ti_le_hop_le", "trang_thai", "goi_y_loai", "tieu_de_goi_y", "ghi_chu"]


def chuan_hoa(s):
    return unicodedata.normalize("NFC", s)


def ti_le_hop_le(text):
    if not text:
        return 0.0
    xau = 0
    for c in text:
        o = ord(c)
        if c == "\ufffd" or 0xE000 <= o <= 0xF8FF or (o < 32 and c not in "\n\r\t"):
            xau += 1
    return 1 - xau / len(text)


def so_tu_pho_bien(text):
    t = chuan_hoa(text).lower()
    return sum(1 for w in TU_PHO_BIEN if w in t)


def goi_y_loai(dong_dau):
    """Chỉ là gợi ý để bạn duyệt, không phải kết luận."""
    for d in dong_dau:
        u = d.upper().strip()
        if re.match(r"^(LUẬT|BỘ LUẬT|NGHỊ ĐỊNH|NGHỊ QUYẾT)\b", u):
            return "nghi_phap_luat"
    for d in dong_dau:
        u = d.upper().strip()
        if re.match(r"^(ĐƠN|GIẤY ĐỀ NGHỊ|PHIẾU|BẢN KHAI|MẪU)\b", u):
            return "nghi_mau_don"
    return ""


def dong_dau_cua(text, n=15):
    return [d.strip() for d in chuan_hoa(text).splitlines() if d.strip()][:n]


def doc_pdf(path):
    import fitz  # pymupdf
    ghi_chu = []
    with fitz.open(path) as doc:
        if doc.needs_pass:
            return None, "pdf_co_mat_khau", ghi_chu
        trang = [p.get_text() for p in doc]
    so_trang = len(trang)
    co_chu = [t for t in trang if len(t.strip()) >= NGUONG_TRANG_CO_CHU]
    text = "\n".join(co_chu)
    return dict(so_trang=so_trang, trang_co_chu=len(co_chu), text=text), None, ghi_chu


def doc_docx(path):
    import docx
    d = docx.Document(str(path))
    text = "\n".join(p.text for p in d.paragraphs)
    for b in d.tables:
        for hang in b.rows:
            text += "\n" + " | ".join(o.text for o in hang.cells)
    return dict(so_trang="", trang_co_chu="", text=text), None, []


def phan_loai(path):
    ext = path.suffix.lower()
    dong = {c: "" for c in COT}
    dong.update(file=path.name, dinh_dang=ext.lstrip("."), kb=int(path.stat().st_size / 1024))
    try:
        if ext == ".pdf":
            kq, loi, _ = doc_pdf(path)
        elif ext == ".docx":
            kq, loi, _ = doc_docx(path)
        elif ext == ".doc":
            dong["trang_thai"] = "can_chuyen_doi"
            dong["ghi_chu"] = "Word cũ: chạy lệnh chuyen-doc"
            return dong
        elif ext == ".txt":
            text = path.read_text(encoding="utf-8", errors="replace")
            kq, loi = dict(so_trang="", trang_co_chu="", text=text), None
        else:
            dong["trang_thai"] = "khong_ho_tro"
            return dong
    except Exception as e:  # file hỏng, thư viện lỗi...
        dong["trang_thai"] = "loi_doc_file"
        dong["ghi_chu"] = f"{type(e).__name__}: {e}"[:200]
        return dong

    if loi:
        dong["trang_thai"] = loi
        return dong

    text = kq["text"]
    dong["so_trang"] = kq["so_trang"]
    dong["trang_co_chu"] = kq["trang_co_chu"]
    dong["so_ky_tu"] = len(text.strip())
    dong["ti_le_hop_le"] = round(ti_le_hop_le(text), 3)
    dau = dong_dau_cua(text)
    dong["goi_y_loai"] = goi_y_loai(dau)
    dong["tieu_de_goi_y"] = " / ".join(dau[:3])[:150]

    # Quyết định trạng thái
    if ext == ".pdf":
        n, c = kq["so_trang"], kq["trang_co_chu"]
        if n == 0 or c == 0:
            dong["trang_thai"] = "can_ocr"
            dong["ghi_chu"] = "Không có trang nào có chữ (scan)"
            return dong
        if c < n:
            dong["ghi_chu"] = f"Chỉ {c}/{n} trang có chữ (scan một phần)"
            dong["trang_thai"] = "scan_mot_phan"
    if len(text.strip()) >= 300 and (
        dong["ti_le_hop_le"] < NGUONG_TI_LE_HOP_LE or so_tu_pho_bien(text) < 3
    ):
        dong["trang_thai"] = "loi_font"
        dong["ghi_chu"] = (dong["ghi_chu"] + "; " if dong["ghi_chu"] else "") + \
            "Có chữ nhưng không ra từ tiếng Việt -> coi như scan, cần OCR/bản gốc"
    elif len(text.strip()) < 200 and ext == ".docx":
        dong["trang_thai"] = "docx_it_chu"
        dong["ghi_chu"] = "Ít chữ, có thể là ảnh dán vào Word"
    elif not dong["trang_thai"]:
        dong["trang_thai"] = "ok"
    return dong


def lenh_quet(a):
    nguon = Path(a.nguon)
    files = sorted(p for p in nguon.rglob("*") if p.is_file())
    if not files:
        print(f"Không thấy file nào trong {nguon.resolve()}")
        return
    Path(a.ra).parent.mkdir(parents=True, exist_ok=True)
    thong_ke = {}
    with open(a.ra, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=COT)
        w.writeheader()
        for i, p in enumerate(files, 1):
            d = phan_loai(p)
            w.writerow(d)
            thong_ke[d["trang_thai"]] = thong_ke.get(d["trang_thai"], 0) + 1
            print(f"[{i}/{len(files)}] {p.name}: {d['trang_thai']}")
    print("\nTổng kết:")
    for k, v in sorted(thong_ke.items(), key=lambda x: -x[1]):
        print(f"  {k}: {v}")
    print(f"\nĐã ghi {a.ra}. Mở bằng Excel, xem các dòng khác 'ok' và cột goi_y_loai.")


def tim_soffice():
    p = shutil.which("soffice") or shutil.which("soffice.exe")
    if p:
        return p
    for c in [r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"]:
        if Path(c).exists():
            return c
    return None


def lenh_chuyen_doc(a):
    soffice = tim_soffice()
    if not soffice:
        print("Không tìm thấy LibreOffice. Cài từ https://www.libreoffice.org/download/ rồi chạy lại.")
        return
    nguon, dich = Path(a.nguon), Path(a.dich)
    dich.mkdir(parents=True, exist_ok=True)
    files = sorted(p for p in nguon.rglob("*.doc") if p.is_file())
    if a.toi_da:
        files = files[: a.toi_da]
    ket_qua = []
    for i, p in enumerate(files, 1):
        out = dich / (p.stem + ".docx")
        trang_thai, ghi_chu = "OK", ""
        try:
            r = subprocess.run(
                [soffice, "--headless", "--convert-to", "docx", "--outdir", str(dich), str(p)],
                capture_output=True, text=True, timeout=120)
            if not out.exists():
                trang_thai, ghi_chu = "LOI", (r.stderr or r.stdout)[:200]
        except subprocess.TimeoutExpired:
            trang_thai, ghi_chu = "LOI", "Quá 120 giây"
        print(f"[{i}/{len(files)}] {p.name}: {trang_thai}")
        ket_qua.append([p.name, out.name, trang_thai, ghi_chu])
    with open(dich / "ket_qua_chuyen_doi.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["file_goc", "file_docx", "trang_thai", "ghi_chu"])
        w.writerows(ket_qua)
    print(f"\nXong. Kết quả ở {dich}. File .docx mới chỉ chép vào raw_documents sau khi bạn kiểm tra có chữ.")
    print("Tip: chạy lại lệnh 'quet' trỏ --nguon vào thư mục đích để kiểm tra các file vừa chuyển.")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="lenh", required=True)
    q = sub.add_parser("quet")
    q.add_argument("--nguon", default=str(NGUON_MAC_DINH))
    q.add_argument("--ra", default=str(RA_MAC_DINH))
    q.set_defaults(ham=lenh_quet)
    c = sub.add_parser("chuyen-doc")
    c.add_argument("--nguon", default=str(NGUON_MAC_DINH))
    c.add_argument("--dich", default=str(DICH_DOC_MAC_DINH))
    c.add_argument("--toi-da", type=int, default=0, help="chỉ chuyển N file đầu để thử")
    c.set_defaults(ham=lenh_chuyen_doc)
    a = ap.parse_args()
    a.ham(a)


if __name__ == "__main__":
    main()
